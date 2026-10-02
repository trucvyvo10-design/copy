import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import os

print("🔄 Running Automated MI Reporting Pipeline...")

# 1. Load Raw Core Banking Data
data_path = 'raw_core_banking_logs.csv'
if not os.path.exists(data_path):
    # Generate dummy data frame if file not found locally
    print("⚠️ Raw log file not found. Creating temporary data in memory...")
    df = pd.DataFrame({
        'transaction_id': ['TXN001', 'TXN002', 'TXN003'],
        'account_number': ['ACC001', 'ACC002', 'ACC003'],
        'transaction_timestamp': ['2026-01-01 10:00:00', '2026-01-01 10:05:00', '2026-01-01 10:10:00'],
        'transaction_type': ['TRANSFER_OUT', 'DEPOSIT', 'TRANSFER_OUT'],
        'amount': [9500.00, 1200.00, 9800.00],
        'currency': ['USD', 'USD', 'USD'],
        'channel': ['MOBILE_BANKING', 'ATM', 'INTERNET_BANKING'],
        'is_suspicious_flag': [1, 0, 1]
    })
else:
    df = pd.read_csv(data_path)

# 2. Process & Aggregate Risk Metrics
summary_by_channel = df.groupby('channel').agg(
    total_transactions=('transaction_id', 'count'),
    total_volume=('amount', 'sum'),
    suspicious_count=('is_suspicious_flag', 'sum')
).reset_index()

# Filter High Risk Structuring Transactions (Amount between $8,000 and $9,999)
high_risk_df = df[df['amount'].between(8000, 9999.99)]

# 3. Export to Formatted Excel Report using openpyxl
output_file = 'Management_Information_Risk_Report.xlsx'
with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
    df.to_excel(writer, sheet_name='Raw_Logs', index=False)
    summary_by_channel.to_excel(writer, sheet_name='MI_Channel_Summary', index=False)
    high_risk_df.to_excel(writer, sheet_name='AML_Structuring_Alerts', index=False)

# 4. Apply Professional Formatting via openpyxl
wb = openpyxl.load_workbook(output_file)
ws = wb['AML_Structuring_Alerts']

# Styling elements
red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid')

# Format Headers
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')

# Highlight High Risk Rows
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
    # Check amount column (column 5)
    if row[4].value and 8000 <= float(row[4].value) <= 9999.99:
        for cell in row:
            cell.fill = red_fill

wb.save(output_file)
print(f"✅ MI Risk Report successfully created and formatted: {output_file}")
